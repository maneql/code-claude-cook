#!/usr/bin/env python3
"""Render a project spec JSON into a styled, multi-sheet .xlsx (developer blueprint).

Usage:  python3 scripts/export-spec.py <spec.json> [out.xlsx]

The script is generic and language-agnostic: it renders WHATEVER sheets,
columns, and rows the JSON defines (no fixed count), and every human-readable
string comes from the JSON, so the output language is whatever the JSON used.
Schema (sheets/columns are arbitrary — shape them per project):

{
  "meta":  {"project","version","language","generated","source"},
  "sheets":[
     {"name": str, "type":"kv",    "rows":[[key,value], ...]},
     {"name": str, "type":"table", "columns":[...], "rows":[[...], ...]}
  ]
}
"""
import sys, json, os, re


def _ensure_openpyxl():
    # Never install packages silently: the user decides what goes into their Python.
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        sys.stderr.write("ERROR: openpyxl is required. Install it for this Python, "
                         "then re-run:\n  \"%s\" -m pip install openpyxl\n" % sys.executable)
        return False


def _put(ws, row, col, value):
    """Write one cell. Text starting with '=' stays text, so a value in the spec
    can never turn into a live Excel formula (formula injection)."""
    cell = ws.cell(row, col, value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def main():
    if len(sys.argv) < 2:
        print("usage: export-spec.py <spec.json> [out.xlsx]"); sys.exit(1)
    src = sys.argv[1]
    with open(src, encoding="utf-8") as fh:
        spec = json.load(fh)
    meta = spec.get("meta", {})
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + ".xlsx"
    if not _ensure_openpyxl():
        sys.exit(2)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(); wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    key_font = Font(bold=True)
    wrap = Alignment(vertical="top", wrap_text=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def autosize(ws, ncols, maxw=70):
        for c in range(1, ncols + 1):
            letter = get_column_letter(c); best = 10
            for cell in ws[letter]:
                v = "" if cell.value is None else str(cell.value)
                longest = max((len(x) for x in v.split("\n")), default=0)
                best = max(best, longest + 2)
            ws.column_dimensions[letter].width = min(best, maxw)

    sheets = spec.get("sheets") or [
        {"name": "Overview", "type": "kv",
         "rows": [["Project", meta.get("project", "")]]}]

    for sh in sheets:
        name = (re.sub(r"[\\/*?:\[\]]", "-", str(sh.get("name") or "Sheet"))[:31]) or "Sheet"
        ws = wb.create_sheet(title=name)
        if sh.get("type") == "kv":
            for r, pair in enumerate(sh.get("rows", []), start=1):
                k = pair[0] if len(pair) > 0 else ""
                v = pair[1] if len(pair) > 1 else ""
                a = _put(ws, r, 1, k); a.font = key_font; a.alignment = wrap; a.border = border
                b = _put(ws, r, 2, v); b.alignment = wrap; b.border = border
            autosize(ws, 2, maxw=100)
        else:
            cols = sh.get("columns", [])
            for c, h in enumerate(cols, start=1):
                cell = _put(ws, 1, c, h)
                cell.font = head_font; cell.fill = head_fill
                cell.alignment = wrap; cell.border = border
            for r, row in enumerate(sh.get("rows", []), start=2):
                for c in range(len(cols)):
                    val = row[c] if c < len(row) else ""
                    if isinstance(val, (list, dict)):
                        val = json.dumps(val, ensure_ascii=False)
                    cell = _put(ws, r, c + 1, val)
                    cell.alignment = wrap; cell.border = border
            if cols:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ws.max_row,1)}"
            autosize(ws, len(cols) or 1)

    wb.properties.title = (f"{meta.get('project','Project')} - Spec "
                           f"{meta.get('version','')}").strip()
    wb.properties.creator = "Claude dev team"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
